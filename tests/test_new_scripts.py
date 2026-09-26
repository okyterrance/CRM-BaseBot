"""2026-09-25 加的两个脚本：登记 AI 客户、ECAS 追加导入。用假 Base 从头到尾跑一遍 main()。

两件最要紧的事：

1. **不改已有的**：已登记的 UID 跳过；ECAS 表里已有的申请一笔不动、不重复加。
2. **预演一个字都不写**。
"""

from __future__ import annotations

import importlib.util
import sys
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from crm_basebot.domain import ecas, schema
from crm_basebot.domain.dates import DEFAULT_BUSINESS_TIMEZONE, date_to_ms
from crm_basebot.lark.bitable import FieldInfo, TableInfo

from .conftest import TBL_AUDIT, TBL_BOARD, TBL_CLIENT, TBL_ECAS, TBL_REFERRAL, TBL_SALES

SGT = DEFAULT_BUSINESS_TIMEZONE
PRANCE = "ou_prance00000000000000000000000"


def _load(name: str):
    path = Path(__file__).resolve().parent.parent / "scripts" / f"{name}.py"
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


register = _load("register_ai_clients")
appender = _load("append_ecas")
relinker = _load("relink_board")


def _settings():
    return SimpleNamespace(
        base_app_token="app",
        table_client=TBL_CLIENT,
        table_referral=TBL_REFERRAL,
        table_sales=TBL_SALES,
        table_audit=TBL_AUDIT,
        table_ecas=TBL_ECAS,
        table_daily_board=TBL_BOARD,
        business_timezone="Asia/Singapore",
    )


def _wire(monkeypatch, module, fake_bitable):
    monkeypatch.setattr(module, "load_settings", _settings)
    monkeypatch.setattr(module, "require_settings", lambda *a, **k: None)
    monkeypatch.setattr(module, "BitableClient", lambda token: fake_bitable)


def _xlsx(tmp_path: Path, header: list, rows: list[list], name="list.xlsx") -> Path:
    book = Workbook()
    sheet = book.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    path = tmp_path / name
    book.save(path)
    return path


@pytest.fixture
def r095(fake_bitable):
    referral = fake_bitable.tables[TBL_REFERRAL].add_existing(
        {
            schema.REFERRAL_NO: "R095",
            schema.REFERRAL_NAME: "JIANG JUN",
            schema.REFERRAL_OWNER_OPEN_ID: PRANCE,
            schema.REFERRAL_SALES_NAME: "Prance Wang",
        }
    )
    fake_bitable.tables[TBL_SALES].add_existing(
        {
            schema.SALES_OPEN_ID: PRANCE,
            schema.SALES_NAME: "Prance Wang",
            schema.SALES_ROLE: schema.ROLE_SALES,
            schema.SALES_STATUS: schema.SALES_STATUS_ACTIVE,
        }
    )
    fake_bitable.tables[TBL_CLIENT].fields = [
        FieldInfo(f"f{i}", name, 1, "Text", False)
        for i, name in enumerate([schema.CLIENT_AI_STATUS, schema.CLIENT_AI_DATE])
    ]
    return referral


# ---------- 登记 AI 客户 ----------

AI_HEADER = ["客户名称", "客户UID", "AI状态", "升级AI日期"]


def _clients(fake_bitable) -> list[dict]:
    return list(fake_bitable.tables[TBL_CLIENT].records.values())


def test_预演一个字都不写(tmp_path, monkeypatch, fake_bitable, r095, capsys):
    _wire(monkeypatch, register, fake_bitable)
    path = _xlsx(
        tmp_path, AI_HEADER, [["WILAI SONNAM", "2219795833498687744", "升级为AI", "2026-08-14"]]
    )

    assert register.main(["--file", str(path), "--referral", "R095"]) == 0
    assert _clients(fake_bitable) == []
    out = capsys.readouterr().out
    assert "登记人：Prance Wang" in out
    assert "要登记 1 个" in out


def test_登记到渠道名下_登记人是渠道负责人(tmp_path, monkeypatch, fake_bitable, r095):
    _wire(monkeypatch, register, fake_bitable)
    path = _xlsx(
        tmp_path,
        AI_HEADER,
        [
            ["WILAI SONNAM", "2219795833498687744", "升级为AI", "2026-08-14"],
            ["SUDARAT SUPHAROEK", "2256003665583491328", "升级为AI", datetime(2026, 9, 1)],
        ],
    )

    assert register.main(["--file", str(path), "--referral", "R095", "--apply"]) == 0

    rows = _clients(fake_bitable)
    assert [row[schema.CLIENT_UID] for row in rows] == [
        "2219795833498687744",
        "2256003665583491328",
    ]
    for row in rows:
        assert row[schema.CLIENT_REFERRAL_LINK] == [r095]
        assert row[schema.CLIENT_OWNER_OPEN_ID] == PRANCE
        assert row[schema.CLIENT_AI_STATUS] == schema.AI_STATUS_UPGRADED
    assert rows[1][schema.CLIENT_AI_DATE] == date_to_ms(date(2026, 9, 1), tz=SGT)


def test_已登记的UID跳过不改(tmp_path, monkeypatch, fake_bitable, r095, capsys):
    fake_bitable.tables[TBL_CLIENT].add_existing(
        {schema.CLIENT_UID: "2171813839481632768", schema.CLIENT_NAME: "XIAOJIA CAI"}
    )
    _wire(monkeypatch, register, fake_bitable)
    path = _xlsx(tmp_path, AI_HEADER, [["CAI XIAOJIA", "2171813839481632768", "开户即AI", None]])

    assert register.main(["--file", str(path), "--referral", "R095", "--apply"]) == 0
    assert len(_clients(fake_bitable)) == 1
    assert "已经登记过（XIAOJIA CAI），跳过不改" in capsys.readouterr().out


def test_数字格子的UID不写(tmp_path, monkeypatch, fake_bitable, r095, capsys):
    _wire(monkeypatch, register, fake_bitable)
    path = _xlsx(tmp_path, AI_HEADER, [["X", 2219795833498687744.0, "开户即AI", None]])

    register.main(["--file", str(path), "--referral", "R095", "--apply"])
    assert _clients(fake_bitable) == []
    assert "数字格子" in capsys.readouterr().out


def test_升级为AI没填日期的不写(tmp_path, monkeypatch, fake_bitable, r095, capsys):
    _wire(monkeypatch, register, fake_bitable)
    path = _xlsx(tmp_path, AI_HEADER, [["X", "2219795833498687744", "升级为AI", None]])

    register.main(["--file", str(path), "--referral", "R095", "--apply"])
    assert _clients(fake_bitable) == []
    assert "要填升级日期" in capsys.readouterr().out


def test_客户表还没有AI两列时停下(tmp_path, monkeypatch, fake_bitable, r095, capsys):
    fake_bitable.tables[TBL_CLIENT].fields = []
    _wire(monkeypatch, register, fake_bitable)
    path = _xlsx(tmp_path, AI_HEADER, [["X", "2219795833498687744", "开户即AI", None]])

    assert register.main(["--file", str(path), "--referral", "R095", "--apply"]) == 1
    assert "sync_base.py --apply" in capsys.readouterr().out


def test_渠道没有负责人时停下(tmp_path, monkeypatch, fake_bitable, r095, capsys):
    fake_bitable.tables[TBL_REFERRAL].records[r095][schema.REFERRAL_OWNER_OPEN_ID] = ""
    _wire(monkeypatch, register, fake_bitable)
    path = _xlsx(tmp_path, AI_HEADER, [["X", "2219795833498687744", "开户即AI", None]])

    assert register.main(["--file", str(path), "--referral", "R095", "--apply"]) == 1
    assert _clients(fake_bitable) == []


# ---------- ECAS 追加 ----------

EXPORT_HEADER = [
    "UID",
    "渠道",
    "币种",
    "eCAS账号",
    "状态",
    "创建时间",
    "审批单号",
    "引用ID",
    "账户名称",
    "ShortName",
    "收费币种",
    "收费金额",
    "销售名称",
]


def _export_row(name, amount, when, sales, ref, uid="2305298134153303296", status="APPROVED"):
    # 审批单号故意都一样：一张审批单能批好几笔，真数据里 XIAOJIA 三笔就共用一个。
    return [
        uid,
        "SC_SG",
        "USD",
        "0179966510",
        status,
        when,
        "202608310102",
        ref,
        name,
        name,
        "USDT",
        amount,
        sales,
    ]


@pytest.fixture
def appending(monkeypatch, fake_bitable, r095):
    _wire(monkeypatch, appender, fake_bitable)
    monkeypatch.setattr(appender, "get_client", lambda: None)
    monkeypatch.setattr(appender.import_ecas, "ensure_table", lambda *a, **k: TBL_ECAS)
    return fake_bitable


def _ecas_rows(fake_bitable) -> list[dict]:
    return list(fake_bitable.tables[TBL_ECAS].records.values())


def _september(tmp_path):
    return _xlsx(
        tmp_path,
        EXPORT_HEADER,
        [
            _export_row(
                "SUNTIPAB SIWONGCHAI", 5000, datetime(2026, 9, 23, 12, 26), "Prance Wang", "A1"
            ),
            _export_row(
                "BLUE ARK LIMITED", 10000, datetime(2026, 9, 24, 17, 20), "Ingrid Lin", "A2"
            ),
            _export_row("HK XIAOJIA", 2000, datetime(2026, 9, 11, 10, 40), "Weichao Huang", "A3"),
            _export_row("HK XIAOJIA", 2000, datetime(2026, 9, 11, 10, 39), "Weichao Huang", "A4"),
            _export_row(
                "PENDING CO", 5000, datetime(2026, 9, 24), "Ingrid Lin", "A5", status="PENDING"
            ),
        ],
        name="export.xlsx",
    )


ASSIGN = ["--assign", "Prance Wang=R095:50"]


def test_ECAS预演一个字都不写(tmp_path, appending, capsys):
    assert appender.main(["--file", str(_september(tmp_path)), *ASSIGN]) == 0
    assert _ecas_rows(appending) == []
    out = capsys.readouterr().out
    assert "4 笔已批准的申请" in out
    assert "R095 JIANG JUN：返佣合计 2,500.00" in out
    assert "状态是 PENDING" in out


def test_追加新申请_认领的记到渠道(tmp_path, appending, r095):
    assert appender.main(["--file", str(_september(tmp_path)), *ASSIGN, "--apply"]) == 0

    rows = {row[ecas.ECAS_REF_ID]: row for row in _ecas_rows(appending)}
    assert set(rows) == {"A1", "A2", "A3", "A4"}, "同名同日同金额的两笔 XIAOJIA 都要进"
    prance = rows["A1"]
    assert prance[ecas.ECAS_REFERRAL_LINK] == [r095]
    assert prance[ecas.ECAS_RATE] == 50.0
    assert prance[ecas.ECAS_REFERRER_NAME] == "JIANG JUN"
    assert prance[ecas.ECAS_CLIENT_UID] == "2305298134153303296"
    assert prance[ecas.ECAS_APPLIED_AT] == int(
        datetime(2026, 9, 23, 12, 26, tzinfo=SGT).timestamp() * 1000
    )
    assert ecas.ECAS_RATE not in rows["A2"], "没认领的先空着，之后在 Base 里补"
    assert ecas.ECAS_FEE not in prance, "公式列不能写"


def test_再跑一次不会重复加(tmp_path, appending):
    path = _september(tmp_path)
    appender.main(["--file", str(path), *ASSIGN, "--apply"])
    appender.main(["--file", str(path), *ASSIGN, "--apply"])
    assert len(_ecas_rows(appending)) == 4


def test_老记录一笔不动_同名同日同金额认成已有(tmp_path, appending, capsys):
    """2026-09-25 之前按财务表导的老记录没有审批单号，按「客户名 + 同一天 + 同金额」认。"""
    appending.tables[TBL_ECAS].add_existing(
        {
            ecas.ECAS_CLIENT_NAME: "suntipab  siwongchai",
            ecas.ECAS_AMOUNT: 5000,
            ecas.ECAS_APPLIED_AT: date_to_ms(date(2026, 9, 23), tz=SGT),
            ecas.ECAS_RATE: 20,
        }
    )
    appender.main(["--file", str(_september(tmp_path)), *ASSIGN, "--apply"])

    rows = _ecas_rows(appending)
    assert len(rows) == 4
    assert rows[0][ecas.ECAS_RATE] == 20, "老记录不改"
    assert "同名、同一天、同金额的记录已经在表里" in capsys.readouterr().out


def test_带引用ID的记录只按引用ID认(tmp_path, appending):
    """同名同日同金额的真新申请不能因为长得像就被挡掉。"""
    appending.tables[TBL_ECAS].add_existing(
        {
            ecas.ECAS_CLIENT_NAME: "HK XIAOJIA",
            ecas.ECAS_AMOUNT: 2000,
            ecas.ECAS_APPLIED_AT: date_to_ms(date(2026, 9, 11), tz=SGT),
            ecas.ECAS_REF_ID: "A3",
        }
    )
    appender.main(["--file", str(_september(tmp_path)), *ASSIGN, "--apply"])
    refs = sorted(row.get(ecas.ECAS_REF_ID) for row in _ecas_rows(appending))
    assert refs == ["A1", "A2", "A3", "A4"]


def test_共用一个审批单号的几笔都进(tmp_path, appending):
    """2026-09-25 预演真数据时撞到的：XIAOJIA 9/11 三笔共用审批单号，按它认会丢两笔。"""
    appender.main(["--file", str(_september(tmp_path)), *ASSIGN, "--apply"])
    assert len(_ecas_rows(appending)) == 4


def test_认领到不存在的渠道时停下(tmp_path, appending):
    assert appender.main(["--file", str(_september(tmp_path)), "--assign", "X=R404:50"]) == 1


@pytest.mark.parametrize("bad", ["Prance Wang", "Prance Wang=R095", "Prance Wang=R095:0"])
def test_assign写法不对时停下(bad):
    with pytest.raises(SystemExit):
        appender.parse_assign([bad])


def test_不是ECAS系统导出时停下(tmp_path, appending):
    path = _xlsx(tmp_path, ["Client Name", "ECAS Revenue"], [["X", 1]])
    with pytest.raises(SystemExit):
        appender.main(["--file", str(path)])


# ---------- 整表导入不再冲掉追加进来的行 ----------


def test_整表导入遇到追加进来的行就拒绝(tmp_path, monkeypatch, appending, capsys):
    importer = appender.import_ecas
    appending.tables[TBL_ECAS].add_existing({ecas.ECAS_CLIENT_NAME: "X", ecas.ECAS_REF_ID: "A1"})
    monkeypatch.setattr(importer, "load_settings", _settings)
    monkeypatch.setattr(importer, "require_settings", lambda *a, **k: None)
    monkeypatch.setattr(importer, "BitableClient", lambda token: appending)
    monkeypatch.setattr(
        appending, "list_tables", lambda: [TableInfo(TBL_ECAS, ecas.TABLE_ECAS_NAME)], raising=False
    )
    path = _xlsx(
        tmp_path,
        [
            "Client Name",
            "ECAS Revenue",
            "Application Time",
            "Sales in Charge",
            "UID",
            "Referrer",
            "%",
            "Amount of Referral Fee",
        ],
        [["Y", "5000", datetime(2026, 8, 1), "Prance Wang", None, None, None, "0"]],
    )
    path = path.rename(tmp_path / "old.xlsx")
    book = __import__("openpyxl").load_workbook(path)
    book.active.title = "ECAS"
    book.save(path)

    assert importer.main(["--file", str(path), "--refresh", "--apply"]) == 1
    assert "--wipe-appended" in capsys.readouterr().err
    assert len(_ecas_rows(appending)) == 1


# ---------- relink_board：客户晚登记，今天就补上他以前的交易 ----------


def test_补挂脚本预演不写_加apply才补(monkeypatch, fake_bitable, capsys):
    _wire(monkeypatch, relinker, fake_bitable)
    client = fake_bitable.tables[TBL_CLIENT].add_existing({schema.CLIENT_UID: "577809207768677761"})
    row = fake_bitable.tables[TBL_BOARD].add_existing(
        {schema.BOARD_CLIENT_UID: "577809207768677761"}
    )
    fake_bitable.tables[TBL_BOARD].add_existing({schema.BOARD_CLIENT_UID: "577809207768677769"})

    assert relinker.main([]) == 0
    assert "会补 1 行（1 个客户）" in capsys.readouterr().out
    assert fake_bitable.updates == []

    assert relinker.main(["--apply"]) == 0
    assert "补好了 1 行" in capsys.readouterr().out
    assert fake_bitable.tables[TBL_BOARD].records[row][schema.BOARD_CLIENT_LINK] == [client]
